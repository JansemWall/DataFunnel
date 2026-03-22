import os
import shutil
import time
from datetime import datetime
from threading import Event
from typing import Callable, Optional

import pandas as pd

from config import (
    ARQUIVO_BASES,
    ARQUIVO_HISTORICO,
    PASTA_ORIGENS,
    PASTA_RESULTADOS,
    carregar_json,
    salvar_json,
)

StatusCallback = Optional[Callable[[str], None]]
ChunkProgressCallback = Optional[Callable[[int, int], None]]


def carregar_bases() -> dict[str, str]:
    return carregar_json(ARQUIVO_BASES, {})

def salvar_bases(bases: dict[str, str]) -> None:
    salvar_json(ARQUIVO_BASES, bases)

def carregar_historico() -> list[dict]:
    return carregar_json(ARQUIVO_HISTORICO, [])

def salvar_historico(historico: list[dict]) -> None:
    salvar_json(ARQUIVO_HISTORICO, historico)

def salvar_nova_base(
    bases_salvas: dict[str, str], nome_arquivo: str, caminho_imutavel: str
) -> tuple[dict[str, str], str]:
    base_nome = nome_arquivo
    contador = 1

    while base_nome in bases_salvas and bases_salvas[base_nome] != caminho_imutavel:
        base_nome = f"{nome_arquivo} ({contador})"
        contador += 1

    bases_salvas[base_nome] = caminho_imutavel
    salvar_bases(bases_salvas)
    return bases_salvas, base_nome

def preparar_base(
    caminho_base_ativa: str,
    is_base_no_cofre: bool,
    bases_salvas: dict[str, str],
    timestamp: str,
    status_callback: StatusCallback = None,
) -> tuple[str, str, bool, dict[str, str]]:
    if is_base_no_cofre:
        caminho_origem_imutavel = caminho_base_ativa
        nome_original = os.path.basename(caminho_origem_imutavel)
        if status_callback:
            status_callback("Lendo base do cache local...")
        return caminho_origem_imutavel, nome_original, True, bases_salvas

    nome_original = os.path.basename(caminho_base_ativa)
    caminho_origem_imutavel = os.path.join(PASTA_ORIGENS, f"origem_{timestamp}.csv")

    if caminho_base_ativa.endswith((".xlsx", ".xls")):
        if status_callback:
            status_callback("Importando Excel para o cofre...")
        df_temp = pd.read_excel(caminho_base_ativa)
        df_temp.to_csv(caminho_origem_imutavel, index=False, encoding="utf-8")
        del df_temp
    else:
        if status_callback:
            status_callback("Armazenando CSV no cofre...")
        shutil.copy2(caminho_base_ativa, caminho_origem_imutavel)

    bases_salvas, _ = salvar_nova_base(bases_salvas, nome_original, caminho_origem_imutavel)
    return caminho_origem_imutavel, nome_original, True, bases_salvas

def _iterador_csv(caminho_origem_imutavel: str, tamanho_lote: int):
    try:
        return pd.read_csv(
            caminho_origem_imutavel,
            encoding="utf-8",
            chunksize=tamanho_lote,
            low_memory=False,
        )
    except UnicodeDecodeError:
        return pd.read_csv(
            caminho_origem_imutavel,
            encoding="latin1",
            sep=";",
            chunksize=tamanho_lote,
            low_memory=False,
        )

def _cabecalho_csv(caminho_origem_imutavel: str) -> pd.Index:
    try:
        return pd.read_csv(caminho_origem_imutavel, nrows=0, encoding="utf-8").columns
    except UnicodeDecodeError:
        return pd.read_csv(caminho_origem_imutavel, nrows=0, encoding="latin1", sep=";").columns

def _total_linhas_csv(caminho_origem_imutavel: str) -> int:
    with open(caminho_origem_imutavel, "rb") as arquivo:
        total_linhas = sum(1 for _ in arquivo)
    return max(total_linhas - 1, 0)

def filtrar_em_lotes(
    caminho_origem_imutavel: str,
    regras_de_busca: list[dict[str, str]],
    tamanho_lote: int = 50000,
    stop_event: Optional[Event] = None,
    status_callback: StatusCallback = None,
    progresso_callback: ChunkProgressCallback = None,
) -> pd.DataFrame:
    resultados: list[pd.DataFrame] = []
    linhas_processadas = 0
    linhas_totais = _total_linhas_csv(caminho_origem_imutavel)
    total_chunks = max((linhas_totais + tamanho_lote - 1) // tamanho_lote, 1)
    chunks_concluidos = 0
    iterador = _iterador_csv(caminho_origem_imutavel, tamanho_lote)

    for chunk in iterador:
        if stop_event and stop_event.is_set():
            raise RuntimeError("PROCESSAMENTO_CANCELADO")

        linhas_processadas += len(chunk)
        if status_callback:
            status_callback(f"Aplicando funil de filtros... ({linhas_processadas} linhas)")

        chunk_filtrado = chunk

        for regra in regras_de_busca:
            if stop_event and stop_event.is_set():
                raise RuntimeError("PROCESSAMENTO_CANCELADO")

            if chunk_filtrado.empty:
                break

            termos_lista = [t.strip() for t in regra["termos"].split(",") if t.strip()]
            padrao_busca = "|".join(termos_lista)

            if regra["colunas"]:
                colunas_alvo = [c.strip() for c in regra["colunas"].split(",") if c.strip()]
                colunas_alvo = [c for c in colunas_alvo if c in chunk_filtrado.columns]

                if colunas_alvo:
                    mascara = (
                        chunk_filtrado[colunas_alvo]
                        .astype(str)
                        .apply(lambda col: col.str.contains(padrao_busca, case=False, na=False))
                        .any(axis=1)
                    )
                else:
                    mascara = pd.Series(False, index=chunk_filtrado.index)
            else:
                mascara = (
                    chunk_filtrado.astype(str)
                    .apply(lambda col: col.str.contains(padrao_busca, case=False, na=False))
                    .any(axis=1)
                )

            chunk_filtrado = chunk_filtrado[mascara]

        if not chunk_filtrado.empty:
            resultados.append(chunk_filtrado)

        chunks_concluidos += 1
        if progresso_callback:
            progresso_callback(chunks_concluidos, total_chunks)

    if resultados:
        return pd.concat(resultados)
    return pd.DataFrame(columns=_cabecalho_csv(caminho_origem_imutavel))

def salvar_resultado_csv(df_filtrado: pd.DataFrame, timestamp: str) -> str:
    caminho_resultado_imutavel = os.path.join(PASTA_RESULTADOS, f"resultado_{timestamp}.csv")
    df_filtrado.to_csv(caminho_resultado_imutavel, index=False, encoding="utf-8")
    return caminho_resultado_imutavel

def exportar_dataframe_excel(df: pd.DataFrame, caminho_exportacao: str) -> None:
    df.to_excel(caminho_exportacao, index=False)

def gerar_nome_exportacao(regras_de_busca: list[dict[str, str]], timestamp: str) -> str:
    primeiro_termo = regras_de_busca[0]["termos"].split(",")[0].strip()[:15]
    return f"Exportacao_Busca_{primeiro_termo}_{timestamp}.xlsx"

def gerar_nome_recuperacao(timestamp: Optional[str] = None) -> str:
    ts = timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"Recuperacao_Historico_{ts}.xlsx"

def formatar_regras_log(regras_de_busca: list[dict[str, str]]) -> str:
    return " -> ".join([f"[{r['colunas'] or 'Todas'}]=({r['termos']})" for r in regras_de_busca])

def registrar_execucao_historico(
    timestamp: str,
    data_hora_formatada: str,
    nome_original: str,
    caminho_origem_imutavel: str,
    caminho_resultado_imutavel: str,
    caminho_exportacao: str,
    regras_de_busca: list[dict[str, str]],
    qtd_linhas: int,
) -> None:
    historico = carregar_historico()
    registro = {
        "id": timestamp,
        "data_hora": data_hora_formatada,
        "arquivo_original": nome_original,
        "caminho_origem_imutavel": caminho_origem_imutavel,
        "caminho_resultado_imutavel": caminho_resultado_imutavel,
        "caminho_exportacao": caminho_exportacao,
        "termos": formatar_regras_log(regras_de_busca),
        "linhas": qtd_linhas,
        "regras_raw": regras_de_busca,
    }
    historico.insert(0, registro)
    salvar_historico(historico)

def registrar_conversao_historico(
    timestamp: str,
    data_hora_formatada: str,
    nome_original: str,
    caminho_csv_gerado: str,
    qtd_linhas: int,
) -> None:
    historico = carregar_historico()
    registro = {
        "id": timestamp,
        "data_hora": data_hora_formatada,
        "arquivo_original": nome_original,
        "caminho_origem_imutavel": caminho_csv_gerado,
        "caminho_resultado_imutavel": caminho_csv_gerado,
        "caminho_exportacao": "",
        "termos": "Conversão Direta (Excel para CSV)",
        "linhas": qtd_linhas,
        "regras_raw": [],
    }
    historico.insert(0, registro)
    salvar_historico(historico)

def atualizar_caminho_exportacao_historico(registro_id: str, novo_caminho: str) -> None:
    historico = carregar_historico()
    for item in historico:
        if item.get("id") == registro_id:
            item["caminho_exportacao"] = novo_caminho
            break
    salvar_historico(historico)

def exportar_resultado_historico(caminho_resultado_imutavel: str, caminho_salvar: str) -> None:
    df = pd.read_csv(caminho_resultado_imutavel, encoding="utf-8")
    df.to_excel(caminho_salvar, index=False)

def gerar_nome_csv_cru(nome_original: str, timestamp: str) -> str:
    nome_base = os.path.splitext(nome_original)[0]
    return f"CSV_Cru_{nome_base}_{timestamp}.csv"

def exportar_csv_cru(
    caminho_origem_imutavel: str,
    caminho_salvar: str,
    status_callback: StatusCallback = None,
) -> None:
    if status_callback:
        status_callback("Preparando exportação do CSV cru...")
    
    try:
        try:
            df = pd.read_csv(caminho_origem_imutavel, encoding="utf-8")
        except UnicodeDecodeError:
            df = pd.read_csv(caminho_origem_imutavel, encoding="latin1", sep=";")
        
        df.to_csv(caminho_salvar, index=False, encoding="utf-8")
        
        if status_callback:
            status_callback(f"CSV cru exportado com sucesso! ({len(df)} linhas)")
    except Exception as erro:
        if status_callback:
            status_callback(f"Erro ao exportar CSV cru: {str(erro)}")
        raise

def _total_linhas_excel(caminho_excel: str) -> int:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(caminho_excel, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        total = max((ws.max_row or 1) - 1, 0)
        wb.close()
        return total
    except Exception:
        df = pd.read_excel(caminho_excel)
        return len(df)

def converter_excel_para_csv_em_lotes(
    caminho_excel: str,
    caminho_csv_destino: str,
    tamanho_lote: int = 2000,
    stop_event: Optional[Event] = None,
    status_callback: StatusCallback = None,
    progresso_callback: ChunkProgressCallback = None,
) -> tuple[int, int]:
    tempo_inicio_preparo = time.perf_counter()
    total_linhas = _total_linhas_excel(caminho_excel)
    
    if total_linhas == 0:
        df_vazio = pd.read_excel(caminho_excel, nrows=0)
        df_vazio.to_csv(caminho_csv_destino, index=False, encoding="utf-8")
        if progresso_callback:
            progresso_callback(1, 1)
        return 0, 1

    total_chunks = max((total_linhas + tamanho_lote - 1) // tamanho_lote, 1)
    etapas_preparo = 4
    total_etapas = total_chunks + etapas_preparo

    def atualizar_preparo(etapa_atual: int, descricao: str) -> None:
        tempo_preparo = time.perf_counter() - tempo_inicio_preparo
        if status_callback:
            status_callback(f"{descricao} ({tempo_preparo:.1f}s)")
        if progresso_callback:
            progresso_callback(etapa_atual, total_etapas)

    atualizar_preparo(1, "Preparando conversao")
    atualizar_preparo(2, "Mapeando linhas da base")
    atualizar_preparo(3, f"Planejando blocos de processamento ({total_chunks} etapas)")

    processados = 0
    chunks_concluidos = 0
    primeiro_chunk = True

    try:
        import openpyxl
        wb = openpyxl.load_workbook(caminho_excel, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        linhas = ws.iter_rows(values_only=True)
        cabecalho = next(linhas, None)
        atualizar_preparo(4, "Estrutura da planilha carregada")

        if not cabecalho:
            pd.DataFrame().to_csv(caminho_csv_destino, index=False, encoding="utf-8")
            wb.close()
            return 0, 1

        buffer_linhas: list[tuple] = []

        for linha in linhas:
            if stop_event and stop_event.is_set():
                raise RuntimeError("PROCESSAMENTO_CANCELADO")

            buffer_linhas.append(linha)

            if len(buffer_linhas) >= tamanho_lote:
                df_chunk = pd.DataFrame(buffer_linhas, columns=list(cabecalho))
                df_chunk.to_csv(
                    caminho_csv_destino,
                    mode="w" if primeiro_chunk else "a",
                    header=primeiro_chunk,
                    index=False,
                    encoding="utf-8",
                )
                processados += len(buffer_linhas)
                chunks_concluidos += 1
                primeiro_chunk = False
                buffer_linhas = []

                if status_callback:
                    status_callback(f"Convertendo base para CSV... ({processados}/{total_linhas} linhas)")
                if progresso_callback:
                    progresso_callback(etapas_preparo + chunks_concluidos, total_etapas)

        if buffer_linhas:
            df_chunk = pd.DataFrame(buffer_linhas, columns=list(cabecalho))
            df_chunk.to_csv(
                caminho_csv_destino,
                mode="w" if primeiro_chunk else "a",
                header=primeiro_chunk,
                index=False,
                encoding="utf-8",
            )
            processados += len(buffer_linhas)
            chunks_concluidos += 1
            if status_callback:
                status_callback(f"Convertendo base para CSV... ({processados}/{total_linhas} linhas)")
            if progresso_callback:
                progresso_callback(etapas_preparo + chunks_concluidos, total_etapas)

        wb.close()

    except ImportError:
        df = pd.read_excel(caminho_excel)
        atualizar_preparo(4, "Estrutura da planilha carregada")
        total_linhas = len(df)
        total_chunks = max((total_linhas + tamanho_lote - 1) // tamanho_lote, 1)
        total_etapas = total_chunks + etapas_preparo
        
        for inicio in range(0, total_linhas, tamanho_lote):
            if stop_event and stop_event.is_set():
                raise RuntimeError("PROCESSAMENTO_CANCELADO")
            
            fim = min(inicio + tamanho_lote, total_linhas)
            chunk = df.iloc[inicio:fim]
            chunk.to_csv(
                caminho_csv_destino,
                mode="w" if inicio == 0 else "a",
                header=inicio == 0,
                index=False,
                encoding="utf-8",
            )
            chunks_concluidos += 1
            processados = fim
            
            if status_callback:
                status_callback(f"Convertendo base para CSV... ({processados}/{total_linhas} linhas)")
            if progresso_callback:
                progresso_callback(etapas_preparo + chunks_concluidos, total_etapas)

    if status_callback:
        status_callback(f"Conversao concluida! ({processados} linhas)")

    return processados, chunks_concluidos